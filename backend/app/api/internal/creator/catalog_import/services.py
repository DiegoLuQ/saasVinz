"""
Carga masiva inicial del catálogo (servicios + productos) de un tenant desde Excel.

La ejecuta el SuperAdmin una sola vez al dar de alta al crematorio:
- No respeta los límites del plan (la hace el admin del SaaS, no el tenant).
- Solo crea: si el servicio (por nombre) o el producto (por código) ya existe,
  la fila se omite y se informa. Nunca actualiza ni borra.
- Todo se lee y escribe bajo el RLS del tenant destino (apply_tenant_rls en el
  router), así que una fila no puede terminar en otro tenant aunque haya un bug aquí.
"""
from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func
from sqlalchemy.orm import Session

from app import models

MAX_ROWS = 2000
SI_NO = ("si", "sí", "no")
EXAMPLE_MARK = "(ejemplo)"  # filas de muestra de la plantilla: se omiten si el admin no las borra

# (encabezado, ancho, obligatorio)
SERVICE_COLUMNS = [
    ("Nombre", 36, True),
    ("Precio venta", 14, True),
    ("Costo", 12, False),
    ("Descripción", 50, False),
    ("Activo (Sí/No)", 14, False),
]
PRODUCT_COLUMNS = [
    ("Código", 14, False),
    ("Nombre", 36, True),
    ("Categoría", 22, False),
    ("Precio venta", 14, True),
    ("Precio costo", 14, False),
    ("Stock", 10, False),
    ("Descripción", 50, False),
    ("Activo (Sí/No)", 14, False),
]


class ImportFileError(ValueError):
    """El archivo no es una plantilla válida (se responde 400)."""


# ---------------------------------------------------------------------------
# Plantilla
# ---------------------------------------------------------------------------

def _write_sheet(ws, columns, example):
    header_fill = PatternFill("solid", fgColor="0F172A")
    for idx, (title, width, required) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=f"{title} *" if required else title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.append(example)
    for cell in ws[2]:
        cell.font = Font(italic=True, color="64748B")
    ws.freeze_panes = "A2"

    active_col = len(columns)
    dv = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    ws.add_data_validation(dv)
    letter = ws.cell(row=1, column=active_col).column_letter
    dv.add(f"{letter}2:{letter}{MAX_ROWS + 1}")


def build_template() -> bytes:
    wb = Workbook()
    ws_info = wb.active
    ws_info.title = "Instrucciones"
    lines = [
        "Carga masiva de catálogo — Vinzer",
        "",
        "1. Completa las hojas 'Servicios' y/o 'Productos'. Puedes dejar una vacía o borrarla y subir solo la otra.",
        "2. Las columnas con * son obligatorias. La fila 2 es un ejemplo: bórrala o reemplázala.",
        "3. Precios en pesos chilenos, sin puntos ni signo $ (ej: 45000).",
        "4. Servicios: si ya existe uno con el mismo nombre, la fila se omite.",
        "5. Productos: si el código ya existe, la fila se omite. Sin código se genera uno (PROD-0001...).",
        "6. Categoría: si no existe en el crematorio, se crea automáticamente.",
        "7. Activo vacío = Sí.",
        f"8. Máximo {MAX_ROWS} filas por hoja. No cambies los nombres de las hojas ni de las columnas.",
    ]
    for line in lines:
        ws_info.append([line])
    ws_info["A1"].font = Font(bold=True, size=14)
    ws_info.column_dimensions["A"].width = 100

    _write_sheet(
        wb.create_sheet("Servicios"),
        SERVICE_COLUMNS,
        ["Cremación individual (ejemplo)", 85000, 30000, "Cremación con devolución de cenizas", "Sí"],
    )
    _write_sheet(
        wb.create_sheet("Productos"),
        PRODUCT_COLUMNS,
        ["URN-001", "Urna de madera (ejemplo)", "Urnas", 35000, 15000, 10, "Urna de pino con grabado", "Sí"],
    )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Lectura / validación
# ---------------------------------------------------------------------------

def _text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).strip()
    return s or None


def _number(value: Any, field: str, errors: List[str], *, required: bool, integer: bool = False):
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            errors.append(f"'{field}' es obligatorio")
        return None
    if isinstance(value, str):
        cleaned = value.strip().replace("$", "").replace(".", "").replace(",", ".").replace(" ", "")
        try:
            value = float(cleaned)
        except ValueError:
            errors.append(f"'{field}' no es un número válido")
            return None
    if not isinstance(value, (int, float)):
        errors.append(f"'{field}' no es un número válido")
        return None
    if value < 0:
        errors.append(f"'{field}' no puede ser negativo")
        return None
    if integer:
        if float(value) != int(value):
            errors.append(f"'{field}' debe ser un número entero")
            return None
        return int(value)
    return float(value)


def _bool(value: Any, errors: List[str]) -> bool:
    s = _text(value)
    if s is None:
        return True
    s = s.lower()
    if s not in SI_NO:
        errors.append("'Activo' debe ser Sí o No")
        return True
    return s != "no"


def _rows(ws, n_cols: int):
    """(número de fila Excel, valores) de las filas con algún dato."""
    for idx, row in enumerate(ws.iter_rows(min_row=2, max_col=n_cols, values_only=True), start=2):
        if all(_text(v) is None for v in row):
            continue
        yield idx, list(row) + [None] * (n_cols - len(row))


def _check_headers(ws, columns):
    expected = [c[0] for c in columns]
    found = [(_text(c.value) or "").rstrip(" *") for c in ws[1][: len(columns)]]
    if found != expected:
        raise ImportFileError(
            f"La hoja '{ws.title}' no tiene las columnas de la plantilla. Descarga la plantilla nuevamente."
        )


def _next_code_factory(existing_codes: set):
    counter = [0]

    def next_code() -> str:
        while True:
            counter[0] += 1
            code = f"PROD-{counter[0]:04d}"
            if code.lower() not in existing_codes:
                existing_codes.add(code.lower())
                return code

    return next_code


def process_import(db: Session, tenant_id: int, content: bytes, *, dry_run: bool) -> Dict[str, Any]:
    """Valida el Excel y, si dry_run es False y no hay errores, crea todo en una transacción."""
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("El archivo no es un Excel válido (.xlsx).")

    # Cada hoja es opcional: se importa lo que venga (Servicios, Productos o ambas).
    # El nombre se compara sin mayúsculas ni espacios de más.
    sheets = {name.strip().lower(): wb[name] for name in wb.sheetnames}
    ws_srv, ws_prd = sheets.get("servicios"), sheets.get("productos")
    if ws_srv is None and ws_prd is None:
        raise ImportFileError("El archivo debe tener una hoja llamada 'Servicios' y/o 'Productos'. Usa la plantilla descargada.")

    srv_rows, prd_rows = [], []
    if ws_srv is not None:
        _check_headers(ws_srv, SERVICE_COLUMNS)
        srv_rows = list(_rows(ws_srv, len(SERVICE_COLUMNS)))
    if ws_prd is not None:
        _check_headers(ws_prd, PRODUCT_COLUMNS)
        prd_rows = list(_rows(ws_prd, len(PRODUCT_COLUMNS)))
    if len(srv_rows) > MAX_ROWS or len(prd_rows) > MAX_ROWS:
        raise ImportFileError(f"Máximo {MAX_ROWS} filas por hoja.")
    if not srv_rows and not prd_rows:
        raise ImportFileError("El archivo no tiene filas para importar.")

    # Datos actuales del tenant (RLS ya aplicado + filtro explícito)
    existing_services = {
        n.lower() for (n,) in db.query(models.Service.name).filter(models.Service.tenant_id == tenant_id) if n
    }
    existing_codes = {
        c.lower() for (c,) in db.query(models.Product.code).filter(models.Product.tenant_id == tenant_id) if c
    }
    categories = {
        c.name.lower(): c
        for c in db.query(models.Category).filter(models.Category.tenant_id == tenant_id)
        if c.name
    }

    errors: List[Dict[str, Any]] = []
    skipped: List[Dict[str, Any]] = []
    services_to_create: List[dict] = []
    products_to_create: List[dict] = []
    new_categories: Dict[str, str] = {}

    # --- Servicios ---
    seen_srv: set = set()
    for row_num, (name, price, cost, desc, active) in srv_rows:
        row_errors: List[str] = []
        name = _text(name)
        if not name:
            row_errors.append("'Nombre' es obligatorio")
        price = _number(price, "Precio venta", row_errors, required=True)
        cost = _number(cost, "Costo", row_errors, required=False)
        is_active = _bool(active, row_errors)
        if row_errors:
            errors.append({"sheet": "Servicios", "row": row_num, "messages": row_errors})
            continue
        key = name.lower()
        if key.endswith(EXAMPLE_MARK):
            skipped.append({"sheet": "Servicios", "row": row_num, "name": name, "reason": "Fila de ejemplo de la plantilla"})
            continue
        if key in seen_srv:
            errors.append({"sheet": "Servicios", "row": row_num, "messages": [f"'{name}' está repetido en el archivo"]})
            continue
        seen_srv.add(key)
        if key in existing_services:
            skipped.append({"sheet": "Servicios", "row": row_num, "name": name, "reason": "Ya existe en el crematorio"})
            continue
        services_to_create.append({
            "name": name, "price": price, "cost": cost or 0.0,
            "description": _text(desc), "is_active": is_active,
        })

    # --- Productos ---
    seen_codes: set = set()
    pending_autocode: List[dict] = []
    for row_num, (code, name, category, price, cost, stock, desc, active) in prd_rows:
        row_errors = []
        code = _text(code)
        name = _text(name)
        if not name:
            row_errors.append("'Nombre' es obligatorio")
        price = _number(price, "Precio venta", row_errors, required=True)
        cost = _number(cost, "Precio costo", row_errors, required=False)
        stock = _number(stock, "Stock", row_errors, required=False, integer=True)
        is_active = _bool(active, row_errors)
        if row_errors:
            errors.append({"sheet": "Productos", "row": row_num, "messages": row_errors})
            continue
        if name.lower().endswith(EXAMPLE_MARK):
            skipped.append({"sheet": "Productos", "row": row_num, "name": name, "reason": "Fila de ejemplo de la plantilla"})
            continue
        if code:
            key = code.lower()
            if key in seen_codes:
                errors.append({"sheet": "Productos", "row": row_num, "messages": [f"El código '{code}' está repetido en el archivo"]})
                continue
            seen_codes.add(key)
            if key in existing_codes:
                skipped.append({"sheet": "Productos", "row": row_num, "name": f"{code} · {name}", "reason": "El código ya existe en el crematorio"})
                continue

        category = _text(category)
        if category and category.lower() not in categories:
            new_categories.setdefault(category.lower(), category)

        item = {
            "code": code, "name": name, "category": category,
            "sale_price": price, "cost_price": cost or 0.0, "stock": stock or 0,
            "description": _text(desc), "is_active": is_active,
        }
        products_to_create.append(item)
        if not code:
            pending_autocode.append(item)

    # Códigos automáticos después de conocer todos los códigos del archivo
    next_code = _next_code_factory(existing_codes | seen_codes)
    for item in pending_autocode:
        item["code"] = next_code()

    summary = {
        "dry_run": dry_run,
        "services": {"to_create": len(services_to_create), "preview": [s["name"] for s in services_to_create[:10]]},
        "products": {"to_create": len(products_to_create), "preview": [f"{p['code']} · {p['name']}" for p in products_to_create[:10]]},
        "categories_to_create": sorted(new_categories.values()),
        "skipped": skipped,
        "errors": errors,
        "applied": False,
    }

    if dry_run or errors or (not services_to_create and not products_to_create):
        return summary

    # --- Escritura (una sola transacción) ---
    try:
        for key, label in new_categories.items():
            cat = models.Category(tenant_id=tenant_id, name=label, description=None)
            db.add(cat)
            categories[key] = cat
        db.flush()

        for s in services_to_create:
            db.add(models.Service(tenant_id=tenant_id, **s))
        for p in products_to_create:
            cat_name = p.pop("category")
            db.add(models.Product(
                tenant_id=tenant_id,
                category_id=categories[cat_name.lower()].id if cat_name else None,
                availability_status="Disponible",  # mismo default que el alta manual
                images=[],
                **p,
            ))
        db.commit()
    except Exception:
        db.rollback()
        raise

    summary["applied"] = True
    return summary
